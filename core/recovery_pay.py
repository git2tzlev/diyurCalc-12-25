"""Recovery pay (דמי הבראה) calculation helpers."""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Optional

import psycopg2.extras

from core.constants import (
    RECOVERY_PAY_DAILY_RATE,
    RECOVERY_PAY_FULL_TIME_MONTHLY_HOURS,
    RECOVERY_PAY_MAX_FTE,
    RECOVERY_PAY_MERAV_CODE,
    RECOVERY_PAY_MONTH,
    TZOHAR_HALEV_HOUSING_ARRAY_ID,
)


LEGACY_2026_MONTHS = {(2025, 7), (2025, 8), (2025, 9), (2025, 10)}
LEGACY_NOTE_RE = re.compile(r"כפול\s*([0-9]+(?:\.[0-9]+)?)\s*משרה")


@dataclass(frozen=True)
class LegacyRecoveryFteRow:
    """A parsed row from the old recovery-pay XLSX file."""

    mifal: str
    meirav_code: str
    id_number: str
    legacy_fte: float
    note: str


def ensure_recovery_pay_legacy_table(conn) -> None:
    """Create the imported old-system FTE table used for 06/2026."""
    cursor = conn.cursor()
    try:
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS recovery_pay_legacy_fte (
                id SERIAL PRIMARY KEY,
                payment_year INTEGER NOT NULL,
                payment_month INTEGER NOT NULL CHECK (payment_month BETWEEN 1 AND 12),
                source_months INTEGER NOT NULL DEFAULT 4,
                person_id INTEGER NULL REFERENCES people(id) ON DELETE SET NULL,
                id_number TEXT NULL,
                meirav_code TEXT NULL,
                mifal TEXT NULL,
                legacy_fte NUMERIC(8, 6) NOT NULL,
                note TEXT NULL,
                created_at TIMESTAMP NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMP NOT NULL DEFAULT NOW(),
                UNIQUE (payment_year, payment_month, id_number, meirav_code)
            )
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_recovery_pay_legacy_period_person
            ON recovery_pay_legacy_fte (payment_year, payment_month, person_id)
        """)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()


def _digits(value: Any) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def _parse_legacy_fte_note(note: Any) -> Optional[float]:
    match = LEGACY_NOTE_RE.search(str(note or ""))
    if not match:
        return None
    return min(float(match.group(1)), RECOVERY_PAY_MAX_FTE)


def parse_legacy_recovery_xlsx(path: str | Path) -> list[LegacyRecoveryFteRow]:
    """Parse the old-system XLSX and extract only the accurate FTE field."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    result: list[LegacyRecoveryFteRow] = []

    for raw in rows:
        row = dict(zip(headers, raw))
        fte = _parse_legacy_fte_note(row.get("Note"))
        if fte is None:
            continue
        result.append(
            LegacyRecoveryFteRow(
                mifal=str(row.get("Mifal") or "").strip(),
                meirav_code=_digits(row.get("Merav")),
                id_number=_digits(row.get("Zehut")),
                legacy_fte=fte,
                note=str(row.get("Note") or "").strip(),
            )
        )
    return result


def import_legacy_recovery_fte(
    conn,
    rows: Iterable[LegacyRecoveryFteRow],
    *,
    payment_year: int = 2026,
    payment_month: int = RECOVERY_PAY_MONTH,
    source_months: int = 4,
) -> int:
    """Upsert parsed old-system FTE rows into the DB table."""
    ensure_recovery_pay_legacy_table(conn)
    parsed_rows = list(rows)
    if not parsed_rows:
        return 0

    people_by_id_number: dict[str, int] = {}
    people_by_meirav: dict[str, int] = {}
    cursor = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute("""
            SELECT id, id_number, meirav_code
            FROM people
            WHERE id_number IS NOT NULL OR meirav_code IS NOT NULL
        """)
        for person in cursor.fetchall():
            id_number = _digits(person["id_number"])
            meirav_code = _digits(person["meirav_code"])
            if id_number:
                people_by_id_number[id_number] = person["id"]
            if meirav_code:
                people_by_meirav[meirav_code] = person["id"]

        for row in parsed_rows:
            person_id = (
                people_by_id_number.get(row.id_number)
                or people_by_meirav.get(row.meirav_code)
            )
            cursor.execute(
                """
                INSERT INTO recovery_pay_legacy_fte
                    (payment_year, payment_month, source_months, person_id,
                     id_number, meirav_code, mifal, legacy_fte, note, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, NOW())
                ON CONFLICT (payment_year, payment_month, id_number, meirav_code)
                DO UPDATE SET
                    source_months = EXCLUDED.source_months,
                    person_id = EXCLUDED.person_id,
                    mifal = EXCLUDED.mifal,
                    legacy_fte = EXCLUDED.legacy_fte,
                    note = EXCLUDED.note,
                    updated_at = NOW()
                """,
                (
                    payment_year,
                    payment_month,
                    source_months,
                    person_id,
                    row.id_number,
                    row.meirav_code,
                    row.mifal,
                    row.legacy_fte,
                    row.note,
                ),
            )
        conn.commit()
        return len(parsed_rows)
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()


def recovery_period_months(payment_year: int) -> list[tuple[int, int]]:
    """Return July previous year through June payment year."""
    return [(payment_year - 1, month) for month in range(7, 13)] + [
        (payment_year, month) for month in range(1, 7)
    ]


def recovery_days_for_seniority(seniority_years: int) -> int:
    """Recovery days by seniority brackets."""
    if seniority_years < 1:
        return 0
    if seniority_years == 1:
        return 5
    if seniority_years <= 3:
        return 6
    if seniority_years <= 10:
        return 7
    if seniority_years <= 15:
        return 8
    if seniority_years <= 19:
        return 9
    return 10


def _seniority_year_for_month(start_date: date, month_start: date) -> int:
    if month_start < start_date:
        return 0
    completed_years = month_start.year - start_date.year
    if (month_start.month, month_start.day) < (start_date.month, start_date.day):
        completed_years -= 1
    return max(1, completed_years + 1)


def recovery_days_for_period(start_date: Any, payment_year: int) -> float:
    """Average recovery days across the July-June recovery year."""
    if not start_date:
        return 0.0
    if hasattr(start_date, "date"):
        start_date = start_date.date()
    if seniority_years_for_recovery(start_date, payment_year) < 1:
        return 0.0

    monthly_days = []
    for year, month in recovery_period_months(payment_year):
        month_start = date(year, month, 1)
        seniority_year = _seniority_year_for_month(start_date, month_start)
        monthly_days.append(recovery_days_for_seniority(seniority_year))
    return round(sum(monthly_days) / len(monthly_days), 2)


def seniority_years_for_recovery(start_date: Any, payment_year: int) -> int:
    """Full recovery-pay years using July 1 as the inclusive yearly boundary."""
    if not start_date:
        return 0
    if hasattr(start_date, "date"):
        start_date = start_date.date()
    cutoff = date(payment_year, RECOVERY_PAY_MONTH + 1, 1)
    years = cutoff.year - start_date.year
    if (cutoff.month, cutoff.day) < (start_date.month, start_date.day):
        years -= 1
    return max(0, years)


def recovery_person_ineligibility_reason(person: Any) -> str:
    """Return a person-level failure without using the current active flag."""
    if not person:
        return "not_found"
    if person["housing_array_id"] != TZOHAR_HALEV_HOUSING_ARRAY_ID:
        return "not_tzohar_halev"
    return ""


def recovery_eligible_minutes_from_totals(
    totals: dict[str, Any],
    *,
    include_standby: bool = False,
) -> int:
    """Eligible minutes for job-scope calculation, without double-counting display splits."""
    minutes = (
        (totals.get("calc100", 0) or 0)
        + (totals.get("calc125", 0) or 0)
        + (totals.get("calc150", 0) or 0)
        + (totals.get("calc175", 0) or 0)
        + (totals.get("calc200", 0) or 0)
        + (totals.get("calc_variable", 0) or 0)
        + (totals.get("vacation_minutes", 0) or 0)
        + (totals.get("sick_minutes", 0) or 0)
    )

    holiday_hours = totals.get("holiday_payment_hours", 0) or 0
    minutes += int(round(float(holiday_hours) * 60))

    if include_standby:
        minutes += int(round(float(totals.get("recovery_standby_minutes", 0) or 0)))

    return int(round(minutes))


def _legacy_hours_for_person(conn, person_id: int, payment_year: int, payment_month: int) -> tuple[float, dict[str, Any]]:
    ensure_recovery_pay_legacy_table(conn)
    cursor = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute(
            """
            SELECT legacy_fte, source_months, id_number, meirav_code, mifal, note
            FROM recovery_pay_legacy_fte
            WHERE payment_year = %s AND payment_month = %s AND person_id = %s
            ORDER BY id DESC
            LIMIT 1
            """,
            (payment_year, payment_month, person_id),
        )
        row = cursor.fetchone()
        if not row:
            return 0.0, {}
        source_months = int(row["source_months"] or 0)
        fte = min(float(row["legacy_fte"] or 0), RECOVERY_PAY_MAX_FTE)
        hours = fte * RECOVERY_PAY_FULL_TIME_MONTHLY_HOURS * source_months
        return hours, dict(row)
    finally:
        cursor.close()


def _has_payment_month_activity(conn, person_id: int, payment_year: int, payment_month: int) -> bool:
    """Require at least one shift or payment component in the payment month."""
    start_date = date(payment_year, payment_month, 1)
    end_date = date(
        payment_year + (1 if payment_month == 12 else 0),
        1 if payment_month == 12 else payment_month + 1,
        1,
    )
    cursor = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute(
            """
            SELECT EXISTS (
                SELECT 1
                FROM time_reports tr
                JOIN apartments ap ON ap.id = tr.apartment_id
                WHERE tr.person_id = %s
                  AND tr.date >= %s
                  AND tr.date < %s
                  AND ap.housing_array_id = %s
            ) OR EXISTS (
                SELECT 1
                FROM payment_components pc
                JOIN apartments ap ON ap.id = pc.apartment_id
                WHERE pc.person_id = %s
                  AND pc.date >= %s
                  AND pc.date < %s
                  AND ap.housing_array_id = %s
            ) AS has_activity
            """,
            (
                person_id, start_date, end_date, TZOHAR_HALEV_HOUSING_ARRAY_ID,
                person_id, start_date, end_date, TZOHAR_HALEV_HOUSING_ARRAY_ID,
            ),
        )
        row = cursor.fetchone()
        return bool(row and row["has_activity"])
    finally:
        cursor.close()


def _inject_holiday_hours(conn, totals: dict[str, Any], person_id: int, year: int, month: int, shabbat_cache: dict, minimum_wage: float, housing_filter: int | None) -> None:
    from core.holiday_payment import calculate_holiday_payments

    hp_map = calculate_holiday_payments(
        conn, year, month, shabbat_cache, minimum_wage,
        housing_filter=housing_filter,
    )
    hp_data = hp_map.get(person_id)
    if not hp_data or hp_data.get("amount", 0) <= 0:
        return
    totals["holiday_payment"] = hp_data["amount"]
    totals["holiday_payment_count"] = hp_data["count"]
    totals["holiday_payment_rate"] = hp_data["rate"]
    details = hp_data.get("details", []) or []
    totals["holiday_payment_details"] = details
    if details:
        totals["holiday_payment_hours"] = round(
            sum(float(item.get("hours") or 0) for item in details),
            2,
        )
    else:
        totals["holiday_payment_hours"] = (
            round(hp_data["amount"] / round(minimum_wage, 2), 2)
            if minimum_wage else 0
        )


def _calculate_month_eligible_minutes(
    conn,
    person_id: int,
    year: int,
    month: int,
    *,
    housing_filter: int | None = None,
    include_standby: bool = False,
) -> int:
    from app_utils import aggregate_daily_segments_to_monthly, get_daily_segments_data
    from core.database import PostgresConnection
    from core.history import get_minimum_wage_for_month
    from core.time_utils import get_shabbat_times_cache

    raw_conn = conn.conn if hasattr(conn, "conn") else conn
    conn_wrapper = conn if hasattr(conn, "execute") else PostgresConnection(raw_conn, use_pool=False)
    minimum_wage = get_minimum_wage_for_month(raw_conn, year, month)
    shabbat_cache = get_shabbat_times_cache(raw_conn)
    daily_segments, _ = get_daily_segments_data(
        conn_wrapper, person_id, year, month, shabbat_cache, minimum_wage
    )
    totals = aggregate_daily_segments_to_monthly(
        conn_wrapper, daily_segments, person_id, year, month, minimum_wage,
        housing_filter=housing_filter,
    )
    _inject_holiday_hours(raw_conn, totals, person_id, year, month, shabbat_cache, minimum_wage, housing_filter)
    return recovery_eligible_minutes_from_totals(totals, include_standby=include_standby)


def calculate_recovery_pay_for_person(
    conn,
    person_id: int,
    payment_year: int,
    payment_month: int,
    *,
    current_month_totals: Optional[dict[str, Any]] = None,
    include_standby: bool = False,
    housing_filter: int | None = None,
) -> dict[str, Any]:
    """Calculate recovery pay for one active guide."""
    if payment_month != RECOVERY_PAY_MONTH:
        return {"amount": 0.0, "eligible": False, "reason": "not_payment_month"}

    raw_conn = conn.conn if hasattr(conn, "conn") else conn
    ensure_recovery_pay_legacy_table(raw_conn)

    cursor = raw_conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute(
            """
            SELECT id, name, start_date, type, housing_array_id
            FROM people
            WHERE id = %s
            """,
            (person_id,),
        )
        person = cursor.fetchone()
    finally:
        cursor.close()

    person_reason = recovery_person_ineligibility_reason(person)
    if person_reason:
        return {"amount": 0.0, "eligible": False, "reason": person_reason}

    seniority_years = seniority_years_for_recovery(person["start_date"], payment_year)
    recovery_days = recovery_days_for_period(person["start_date"], payment_year)
    if recovery_days <= 0:
        return {
            "amount": 0.0,
            "eligible": False,
            "reason": "seniority_under_one_year",
            "seniority_years": seniority_years,
        }
    if not _has_payment_month_activity(raw_conn, person_id, payment_year, payment_month):
        return {
            "amount": 0.0,
            "eligible": False,
            "reason": "no_payment_month_activity",
            "seniority_years": seniority_years,
            "recovery_days": recovery_days,
        }

    legacy_hours, legacy_row = (0.0, {})
    if payment_year == 2026:
        legacy_hours, legacy_row = _legacy_hours_for_person(raw_conn, person_id, payment_year, payment_month)

    eligible_minutes = 0
    calculated_months: list[str] = []
    for year, month in recovery_period_months(payment_year):
        if payment_year == 2026 and (year, month) in LEGACY_2026_MONTHS:
            continue
        if year == payment_year and month == payment_month and current_month_totals is not None:
            eligible_minutes += recovery_eligible_minutes_from_totals(
                current_month_totals,
                include_standby=include_standby,
            )
        else:
            eligible_minutes += _calculate_month_eligible_minutes(
                conn,
                person_id,
                year,
                month,
                housing_filter=housing_filter,
                include_standby=include_standby,
            )
        calculated_months.append(f"{month:02d}/{year}")

    system_hours = eligible_minutes / 60
    total_hours = legacy_hours + system_hours
    denominator = 12 * RECOVERY_PAY_FULL_TIME_MONTHLY_HOURS
    fte = min(RECOVERY_PAY_MAX_FTE, total_hours / denominator) if denominator else 0.0
    amount = round(recovery_days * RECOVERY_PAY_DAILY_RATE * fte, 2)

    return {
        "amount": amount,
        "eligible": amount > 0,
        "reason": "",
        "merav_code": RECOVERY_PAY_MERAV_CODE,
        "daily_rate": RECOVERY_PAY_DAILY_RATE,
        "recovery_days": recovery_days,
        "seniority_years": seniority_years,
        "fte": round(fte, 6),
        "fte_percent": round(fte * 100, 2),
        "legacy_hours": round(legacy_hours, 2),
        "system_hours": round(system_hours, 2),
        "total_hours": round(total_hours, 2),
        "fte_denominator_hours": denominator,
        "calculated_months": calculated_months,
        "legacy_source": legacy_row,
        "include_standby": include_standby,
    }


def apply_recovery_pay_to_totals(totals: dict[str, Any], recovery_data: dict[str, Any]) -> None:
    """Inject calculated recovery pay into monthly totals."""
    amount = recovery_data.get("amount", 0) or 0
    if amount <= 0:
        return
    totals["recovery_pay"] = amount
    totals["recovery_pay_details"] = recovery_data
    rounded = round(round(amount, 2), 1)
    totals["total_payment"] = totals.get("total_payment", 0) + rounded
    totals["gesher_total"] = totals.get("gesher_total", 0) + rounded
    totals["display_total"] = totals.get("display_total", 0) + rounded
    totals["rounded_total"] = totals.get("rounded_total", 0) + rounded
