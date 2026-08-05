"""Annual clothing pay (דמי ביגוד) calculation helpers."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Optional

import psycopg2.extras

from core.constants import (
    CLOTHING_PAY_FULL_TIME_AMOUNT,
    CLOTHING_PAY_FULL_TIME_MONTHLY_HOURS,
    CLOTHING_PAY_MAX_FTE,
    CLOTHING_PAY_MERAV_CODE,
    CLOTHING_PAY_MONTH,
    PERMANENT_EMPLOYEE_TYPE,
    TZOHAR_HALEV_HOUSING_ARRAY_ID,
)
from core.recovery_pay import (
    _calculate_month_eligible_minutes,
    _has_payment_month_activity,
    recovery_eligible_minutes_from_totals,
)


LEGACY_2026_MONTHS = {(2025, 8), (2025, 9), (2025, 10)}


@dataclass(frozen=True)
class LegacyClothingHoursRow:
    """A monthly guide-hours row from the old system."""

    source_year: int
    source_month: int
    id_number: str
    meirav_code: str
    full_name: str
    source_hours: float
    mifal: str


def ensure_clothing_pay_legacy_table(conn) -> None:
    """Create the imported old-system monthly-hours table."""
    cursor = conn.cursor()
    try:
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS clothing_pay_legacy_hours (
                id SERIAL PRIMARY KEY,
                payment_year INTEGER NOT NULL,
                payment_month INTEGER NOT NULL CHECK (payment_month BETWEEN 1 AND 12),
                source_year INTEGER NOT NULL,
                source_month INTEGER NOT NULL CHECK (source_month BETWEEN 1 AND 12),
                person_id INTEGER NULL REFERENCES people(id) ON DELETE SET NULL,
                id_number TEXT NULL,
                meirav_code TEXT NULL,
                full_name TEXT NULL,
                mifal TEXT NULL,
                source_hours NUMERIC(10, 2) NOT NULL,
                created_at TIMESTAMP NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMP NOT NULL DEFAULT NOW(),
                UNIQUE (payment_year, payment_month, source_year, source_month, id_number, meirav_code)
            )
        """)
        cursor.execute("""
            CREATE INDEX IF NOT EXISTS idx_clothing_pay_legacy_period_person
            ON clothing_pay_legacy_hours (payment_year, payment_month, person_id)
        """)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()


def _digits(value: Any) -> str:
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def parse_legacy_clothing_xlsx(path: str | Path) -> list[LegacyClothingHoursRow]:
    """Parse monthly guide hours exported from the old system."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    result: list[LegacyClothingHoursRow] = []
    for raw in rows:
        row = dict(zip(headers, raw))
        source_date = row.get("Hodesh")
        if not source_date or row.get("SHaot") is None:
            continue
        source_hours = max(0.0, float(row["SHaot"] or 0))
        result.append(LegacyClothingHoursRow(
            source_year=int(source_date.year),
            source_month=int(source_date.month),
            id_number=_digits(row.get("Worker")),
            meirav_code=_digits(row.get("Merav")),
            full_name=str(row.get("FullName") or "").strip(),
            source_hours=source_hours,
            mifal=str(row.get("Mifal") or "").strip(),
        ))
    return result


def import_legacy_clothing_hours(
    conn,
    rows: Iterable[LegacyClothingHoursRow],
    *,
    payment_year: int = 2026,
    payment_month: int = CLOTHING_PAY_MONTH,
) -> dict[str, int]:
    """Upsert old-system hours and report matching statistics."""
    ensure_clothing_pay_legacy_table(conn)
    parsed_rows = list(rows)
    stats = {"read": len(parsed_rows), "matched": 0, "unmatched": 0}
    if not parsed_rows:
        return stats

    by_id_number: dict[str, int] = {}
    by_meirav: dict[str, int] = {}
    cursor = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute("""
            SELECT id, id_number, meirav_code FROM people
            WHERE id_number IS NOT NULL OR meirav_code IS NOT NULL
        """)
        for person in cursor.fetchall():
            id_number = _digits(person["id_number"])
            meirav_code = _digits(person["meirav_code"])
            if id_number:
                by_id_number[id_number] = person["id"]
            if meirav_code:
                by_meirav[meirav_code] = person["id"]

        for row in parsed_rows:
            person_id = by_id_number.get(row.id_number) or by_meirav.get(row.meirav_code)
            stats["matched" if person_id else "unmatched"] += 1
            cursor.execute("""
                INSERT INTO clothing_pay_legacy_hours
                    (payment_year, payment_month, source_year, source_month, person_id,
                     id_number, meirav_code, full_name, mifal, source_hours, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW())
                ON CONFLICT (payment_year, payment_month, source_year, source_month, id_number, meirav_code)
                DO UPDATE SET person_id = EXCLUDED.person_id,
                    full_name = EXCLUDED.full_name, mifal = EXCLUDED.mifal,
                    source_hours = EXCLUDED.source_hours, updated_at = NOW()
            """, (
                payment_year, payment_month, row.source_year, row.source_month,
                person_id, row.id_number, row.meirav_code, row.full_name,
                row.mifal, row.source_hours,
            ))
        conn.commit()
        return stats
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()


def clothing_period_months(payment_year: int) -> list[tuple[int, int]]:
    """Return August previous year through July payment year."""
    return [(payment_year - 1, month) for month in range(8, 13)] + [
        (payment_year, month) for month in range(1, 8)
    ]


def has_clothing_seniority(start_date: Any, payment_year: int) -> bool:
    """Use August 1 after payment as the inclusive one-year boundary."""
    if not start_date:
        return False
    if hasattr(start_date, "date"):
        start_date = start_date.date()
    return start_date <= date(payment_year - 1, CLOTHING_PAY_MONTH + 1, 1)


def clothing_fte_from_monthly_hours(monthly_hours: Iterable[float]) -> float:
    """Calculate annual FTE after applying the 182-hour cap per month."""
    capped_hours = sum(
        min(float(CLOTHING_PAY_FULL_TIME_MONTHLY_HOURS), max(0.0, float(hours or 0)))
        for hours in monthly_hours
    )
    denominator = 12 * CLOTHING_PAY_FULL_TIME_MONTHLY_HOURS
    return min(CLOTHING_PAY_MAX_FTE, capped_hours / denominator) if denominator else 0.0


def clothing_person_ineligibility_reason(person: Any, payment_year: int) -> str:
    """Return the first person-level eligibility failure, or an empty string."""
    if not person:
        return "not_found"
    if person["type"] != PERMANENT_EMPLOYEE_TYPE:
        return "not_permanent"
    if person["housing_array_id"] != TZOHAR_HALEV_HOUSING_ARRAY_ID:
        return "not_tzohar_halev"
    if not has_clothing_seniority(person["start_date"], payment_year):
        return "seniority_under_one_year"
    return ""


def _legacy_hours_by_month(conn, person_id: int, payment_year: int, payment_month: int) -> dict[tuple[int, int], float]:
    ensure_clothing_pay_legacy_table(conn)
    cursor = conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute("""
            SELECT source_year, source_month, source_hours
            FROM clothing_pay_legacy_hours
            WHERE payment_year = %s AND payment_month = %s AND person_id = %s
        """, (payment_year, payment_month, person_id))
        return {
            (int(row["source_year"]), int(row["source_month"])): float(row["source_hours"] or 0)
            for row in cursor.fetchall()
        }
    finally:
        cursor.close()


def calculate_clothing_pay_for_person(
    conn,
    person_id: int,
    payment_year: int,
    payment_month: int,
    *,
    current_month_totals: Optional[dict[str, Any]] = None,
    housing_filter: int | None = None,
) -> dict[str, Any]:
    """Calculate annual clothing pay for one eligible permanent guide."""
    if payment_month != CLOTHING_PAY_MONTH:
        return {"amount": 0.0, "eligible": False, "reason": "not_payment_month"}

    raw_conn = conn.conn if hasattr(conn, "conn") else conn
    ensure_clothing_pay_legacy_table(raw_conn)
    cursor = raw_conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
    try:
        cursor.execute("""
            SELECT id, name, start_date, type, housing_array_id
            FROM people WHERE id = %s
        """, (person_id,))
        person = cursor.fetchone()
    finally:
        cursor.close()

    person_reason = clothing_person_ineligibility_reason(person, payment_year)
    if person_reason:
        return {"amount": 0.0, "eligible": False, "reason": person_reason}
    if not _has_payment_month_activity(raw_conn, person_id, payment_year, payment_month):
        return {"amount": 0.0, "eligible": False, "reason": "no_payment_month_activity"}

    legacy_by_month = (
        _legacy_hours_by_month(raw_conn, person_id, payment_year, payment_month)
        if payment_year == 2026 else {}
    )
    capped_hours = 0.0
    legacy_hours = 0.0
    system_hours = 0.0
    month_details: list[dict[str, Any]] = []
    for year, month in clothing_period_months(payment_year):
        if payment_year == 2026 and (year, month) in LEGACY_2026_MONTHS:
            hours = legacy_by_month.get((year, month), 0.0)
            source = "legacy"
            legacy_hours += hours
        elif year == payment_year and month == payment_month and current_month_totals is not None:
            hours = recovery_eligible_minutes_from_totals(current_month_totals) / 60
            source = "system"
            system_hours += hours
        else:
            hours = _calculate_month_eligible_minutes(
                conn, person_id, year, month, housing_filter=housing_filter
            ) / 60
            source = "system"
            system_hours += hours
        month_capped = min(float(CLOTHING_PAY_FULL_TIME_MONTHLY_HOURS), max(0.0, hours))
        capped_hours += month_capped
        month_details.append({
            "year": year, "month": month, "source": source,
            "hours": round(hours, 2), "capped_hours": round(month_capped, 2),
        })

    denominator = 12 * CLOTHING_PAY_FULL_TIME_MONTHLY_HOURS
    fte = clothing_fte_from_monthly_hours(
        item["hours"] for item in month_details
    )
    amount = round(CLOTHING_PAY_FULL_TIME_AMOUNT * fte, 2)
    return {
        "amount": amount,
        "eligible": amount > 0,
        "reason": "",
        "merav_code": CLOTHING_PAY_MERAV_CODE,
        "full_time_amount": CLOTHING_PAY_FULL_TIME_AMOUNT,
        "fte": round(fte, 6),
        "fte_percent": round(fte * 100, 2),
        "legacy_hours": round(legacy_hours, 2),
        "system_hours": round(system_hours, 2),
        "capped_hours": round(capped_hours, 2),
        "fte_denominator_hours": denominator,
        "month_details": month_details,
    }


def apply_clothing_pay_to_totals(totals: dict[str, Any], clothing_data: dict[str, Any]) -> None:
    """Inject calculated clothing pay into monthly totals."""
    amount = clothing_data.get("amount", 0) or 0
    if amount <= 0:
        return
    totals["clothing_pay"] = amount
    totals["clothing_pay_details"] = clothing_data
    rounded = round(round(amount, 2), 1)
    for key in ("total_payment", "gesher_total", "display_total", "rounded_total"):
        totals[key] = totals.get(key, 0) + rounded
